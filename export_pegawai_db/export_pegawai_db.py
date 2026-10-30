import mysql.connector
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_pegawai_to_excel():
    db_config = {
        'host': 'localhost',
        'user': 'rudy',
        'password': '',
        'database': 'dinkes_web'
    }

    try:
        conn = mysql.connector.connect(**db_config)
        cursor = conn.cursor()

        # Query using COALESCE to default to 'Dinas Kesehatan' when u.label is NULL
        query = """
            SELECT 
                p.nip, 
                p.nama,
                COALESCE(u.label, 'Dinas Kesehatan') AS nama_instansi
            FROM pegawai p
            LEFT JOIN unitpelaksanateknis_tbl u ON p.uuid_instansi = u.uuid
            ORDER BY p.nama ASC;
        """
        cursor.execute(query)
        rows = cursor.fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data Pegawai"

        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:E1")
        title_cell = ws["A1"]
        title_cell.value = "DAFTAR NIP, NAMA, DAN INSTANSI PEGAWAI"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 30

        headers = ["No", "NIP", "Nama Pegawai", "Nama Instansi"]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=3, column=col_idx, value=header)

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        ws.row_dimensions[3].height = 24
        for col_num in range(1, 5):
            cell = ws.cell(row=3, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="left" if col_num in (3, 5) else "center", 
                vertical="center"
            )
            cell.border = thin_border

        zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10, color="333333")

        start_row = 4
        for idx, (nip, nama, nama_instansi) in enumerate(rows, start=1):
            current_row = start_row + idx - 1
            ws.row_dimensions[current_row].height = 20

            cell_no = ws.cell(row=current_row, column=1, value=idx)
            cell_nip = ws.cell(row=current_row, column=2, value=str(nip) if nip else "-")
            cell_nama = ws.cell(row=current_row, column=3, value=nama)
            cell_instansi = ws.cell(row=current_row, column=4, value=nama_instansi)

            cell_no.alignment = Alignment(horizontal="center", vertical="center")
            cell_nip.alignment = Alignment(horizontal="center", vertical="center")
            cell_nama.alignment = Alignment(horizontal="left", vertical="center")
            cell_instansi.alignment = Alignment(horizontal="left", vertical="center")

            cell_nip.number_format = '@'

            for cell in (cell_no, cell_nip, cell_nama, cell_instansi):
                cell.font = data_font
                cell.border = thin_border
                if idx % 2 == 0:
                    cell.fill = zebra_fill

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < 3:
                    continue
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        ws.freeze_panes = "A4"

        output_filename = "data_pegawai.xlsx"
        wb.save(output_filename)
        print(f"Data successfully exported to {output_filename}")

    except mysql.connector.Error as err:
        print(f"Database Error: {err}")
    except Exception as e:
        print(f"An error occurred: {e}")
    finally:
        if 'cursor' in locals():
            cursor.close()
        if 'conn' in locals() and conn.is_connected():
            conn.close()

if __name__ == "__main__":
    export_pegawai_to_excel()